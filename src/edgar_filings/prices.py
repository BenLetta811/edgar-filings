"""Import Tableau/Excel stock quote snapshots into SQLite."""

from __future__ import annotations

from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import requests
from openpyxl import load_workbook

from edgar_filings.client import resolve_user_agent
from edgar_filings.db import Database

COLUMNS = (
    "Ticker",
    "Name",
    "Price",
    "Change",
    "Change %",
    "Volume (M)",
    "Avg Vol (3M) (M)",
    "Market Cap ($B)",
    "P/E Ratio (TTM)",
    "52 Wk Change %",
    "52 Wk Low",
    "52 Wk High",
)


def _num(value) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "").replace("%", "").strip())
    except ValueError:
        return None


def _as_of_from_path(path: Path, explicit: str | None) -> str:
    if explicit:
        return datetime.strptime(explicit, "%Y-%m-%d").date().isoformat()
    stamp = datetime.fromtimestamp(path.stat().st_mtime).date()
    return stamp.isoformat()


def parse_stock_prices_xlsx(path: str | Path, as_of: str | None = None) -> list[tuple]:
    file_path = Path(path).expanduser()
    if not file_path.is_file():
        raise FileNotFoundError(f"Excel file not found: {file_path}")
    snapshot = _as_of_from_path(file_path, as_of)
    wb = load_workbook(file_path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            raise ValueError("Workbook has no header row")
        names = [str(col).strip() if col is not None else "" for col in header]
        index = {name: i for i, name in enumerate(names)}
        missing = [col for col in COLUMNS if col not in index]
        if missing:
            raise ValueError(f"Missing columns {missing}; found {names}")

        def cell(row, col):
            i = index[col]
            return row[i] if i < len(row) else None

        out: list[tuple] = []
        for row in rows:
            ticker = str(cell(row, "Ticker") or "").upper().strip()
            if not ticker:
                continue
            out.append(
                (
                    ticker,
                    str(cell(row, "Name") or "").strip(),
                    _num(cell(row, "Price")),
                    _num(cell(row, "Change")),
                    _num(cell(row, "Change %")),
                    _num(cell(row, "Volume (M)")),
                    _num(cell(row, "Avg Vol (3M) (M)")),
                    _num(cell(row, "Market Cap ($B)")),
                    _num(cell(row, "P/E Ratio (TTM)")),
                    _num(cell(row, "52 Wk Change %")),
                    _num(cell(row, "52 Wk Low")),
                    _num(cell(row, "52 Wk High")),
                    snapshot,
                )
            )
        return out
    finally:
        wb.close()


def import_stock_prices(
    db: Database,
    path: str | Path,
    as_of: str | None = None,
) -> tuple[int, str]:
    rows = parse_stock_prices_xlsx(path, as_of=as_of)
    count = db.upsert_quotes(rows)
    snapshot = rows[0][-1] if rows else date.today().isoformat()
    return count, snapshot


CHART_URL = "https://query1.finance.yahoo.com/v8/finance/chart/{ticker}?range=5d&interval=1d"


def fetch_live_quote(ticker: str, session: requests.Session | None = None) -> tuple | None:
    http = session or requests.Session()
    response = http.get(
        CHART_URL.format(ticker=ticker),
        headers={"User-Agent": resolve_user_agent(), "Accept": "application/json"},
        timeout=20,
    )
    if response.status_code != 200:
        return None
    payload = response.json()
    result = ((payload.get("chart") or {}).get("result") or [None])[0]
    if not result:
        return None
    meta = result.get("meta") or {}
    price = meta.get("regularMarketPrice")
    if price is None:
        return None
    prev = meta.get("chartPreviousClose") or meta.get("previousClose")
    change = float(price) - float(prev) if prev is not None else None
    change_pct = (change / float(prev) * 100) if change is not None and prev else None
    volume = meta.get("regularMarketVolume")
    low = meta.get("fiftyTwoWeekLow")
    high = meta.get("fiftyTwoWeekHigh")
    name = str(meta.get("longName") or meta.get("shortName") or ticker).strip()
    return (
        ticker.upper(),
        name,
        float(price),
        change,
        change_pct,
        float(volume) / 1_000_000 if volume is not None else None,
        None,
        None,
        None,
        None,
        float(low) if low is not None else None,
        float(high) if high is not None else None,
    )


def refresh_live_quotes(db: Database, as_of: str | None = None) -> tuple[int, int, str]:
    snapshot = as_of or date.today().isoformat()
    existing = db.latest_quotes(limit=500)
    if not existing:
        raise RuntimeError("No tickers in stock_quotes. Import the Excel file first.")
    session = requests.Session()
    rows = []
    failed = 0
    for quote in existing:
        ticker = str(quote["ticker"])
        live = fetch_live_quote(ticker, session=session)
        if live is None:
            failed += 1
            continue
        rows.append((*live, snapshot))
    count = db.upsert_quotes(rows) if rows else 0
    return count, failed, snapshot


HISTORY_URL = "https://query1.finance.yahoo.com/v8/finance/chart/{ticker}?range=1mo&interval=1d"


def fetch_daily_history(
    ticker: str,
    start: date,
    session: requests.Session | None = None,
    name: str = "",
) -> list[tuple]:
    http = session or requests.Session()
    response = http.get(
        HISTORY_URL.format(ticker=ticker),
        headers={"User-Agent": resolve_user_agent(), "Accept": "application/json"},
        timeout=20,
    )
    if response.status_code != 200:
        return []
    payload = response.json()
    result = ((payload.get("chart") or {}).get("result") or [None])[0]
    if not result:
        return []
    meta = result.get("meta") or {}
    display_name = str(name or meta.get("longName") or meta.get("shortName") or ticker).strip()
    timestamps = result.get("timestamp") or []
    quote = ((result.get("indicators") or {}).get("quote") or [{}])[0]
    closes = quote.get("close") or []
    volumes = quote.get("volume") or []
    low52 = meta.get("fiftyTwoWeekLow")
    high52 = meta.get("fiftyTwoWeekHigh")
    rows: list[tuple] = []
    prev_close = None
    for i, ts in enumerate(timestamps):
        day = datetime.fromtimestamp(int(ts), tz=timezone.utc).date()
        if day < start:
            if i < len(closes) and closes[i] is not None:
                prev_close = float(closes[i])
            continue
        if i >= len(closes) or closes[i] is None:
            continue
        price = float(closes[i])
        change = price - prev_close if prev_close is not None else None
        change_pct = (change / prev_close * 100) if change is not None and prev_close else None
        volume = volumes[i] if i < len(volumes) else None
        rows.append(
            (
                ticker.upper(),
                display_name,
                price,
                change,
                change_pct,
                float(volume) / 1_000_000 if volume is not None else None,
                None,
                None,
                None,
                None,
                float(low52) if low52 is not None else None,
                float(high52) if high52 is not None else None,
                day.isoformat(),
            )
        )
        prev_close = price
    return rows


def load_price_history(db: Database, days: int = 10) -> tuple[int, int, str, str]:
    end = date.today()
    start = end - timedelta(days=max(days, 1) - 1)
    existing = db.latest_quotes(limit=500)
    if not existing:
        raise RuntimeError("No tickers in stock_quotes. Import the Excel file first.")
    session = requests.Session()
    rows: list[tuple] = []
    failed = 0
    for quote in existing:
        ticker = str(quote["ticker"])
        name = str(quote["name"] or "")
        history = fetch_daily_history(ticker, start, session=session, name=name)
        if not history:
            failed += 1
            continue
        rows.extend(history)
    count = db.upsert_quotes(rows) if rows else 0
    return count, failed, start.isoformat(), end.isoformat()


YTD_URL = "https://query1.finance.yahoo.com/v8/finance/chart/{ticker}?range=ytd&interval=1d"


def fetch_ytd(ticker: str, session: requests.Session | None = None) -> dict | None:
    http = session or requests.Session()
    response = http.get(
        YTD_URL.format(ticker=ticker),
        headers={"User-Agent": resolve_user_agent(), "Accept": "application/json"},
        timeout=20,
    )
    if response.status_code != 200:
        return None
    payload = response.json()
    result = ((payload.get("chart") or {}).get("result") or [None])[0]
    if not result:
        return None
    timestamps = result.get("timestamp") or []
    closes = ((result.get("indicators") or {}).get("quote") or [{}])[0].get("close") or []
    pairs = [
        (
            datetime.fromtimestamp(int(ts), tz=timezone.utc).date(),
            float(closes[i]),
        )
        for i, ts in enumerate(timestamps)
        if i < len(closes) and closes[i] is not None
    ]
    year = date.today().year
    pairs = [(day, price) for day, price in pairs if day.year == year]
    if len(pairs) < 1:
        return None
    start_date, start_price = pairs[0]
    last_date, last_price = pairs[-1]
    if start_price == 0:
        return None
    return {
        "ticker": ticker.upper(),
        "year": year,
        "start_date": start_date.isoformat(),
        "start_price": start_price,
        "last_date": last_date.isoformat(),
        "last_price": last_price,
        "ytd_return_pct": (last_price / start_price - 1) * 100,
        "name": str((result.get("meta") or {}).get("longName") or ticker),
    }


def load_ytd_returns(db: Database) -> tuple[int, int]:
    existing = db.latest_quotes(limit=500)
    if not existing:
        raise RuntimeError("No tickers in stock_quotes. Import prices first.")
    session = requests.Session()
    ytd_rows = []
    quote_rows = []
    failed = 0
    for quote in existing:
        ticker = str(quote["ticker"])
        data = fetch_ytd(ticker, session=session)
        if data is None:
            failed += 1
            continue
        ytd_rows.append(
            (
                data["ticker"],
                data["year"],
                data["start_date"],
                data["start_price"],
                data["last_date"],
                data["last_price"],
                data["ytd_return_pct"],
            )
        )
        quote_rows.append(
            (
                data["ticker"],
                str(quote["name"] or data["name"]),
                data["start_price"],
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                data["start_date"],
            )
        )
    db.upsert_quotes(quote_rows)
    db.upsert_ytd(ytd_rows)
    return len(ytd_rows), failed
